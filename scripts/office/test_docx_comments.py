#!/usr/bin/env python3
"""Regression checks for native comment ranges spanning formatted runs."""
import tempfile
import unittest
from pathlib import Path

from docx import Document

from docx_comments import add_comment, edit_comment, list_comments


class NativeCommentsTest(unittest.TestCase):
    def test_formatted_run_boundary_and_table(self):
        with tempfile.TemporaryDirectory() as directory:
            filename = Path(directory) / "comments.docx"
            doc = Document()
            para = doc.add_paragraph()
            para.add_run("前甲").bold = True
            para.add_run("乙丙后")
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "表格甲乙丙"
            first = add_comment(doc, {"quote": "甲乙丙", "text": "删掉乙", "author": "测试"})
            second = add_comment(doc, {"quote": "表格甲乙丙", "text": "检查表格"})
            doc.save(filename)

            reopened = Document(filename)
            comments = {c["id"]: c for c in list_comments(reopened)}
            self.assertEqual(comments[first["id"]]["anchorQuote"], "甲乙丙")
            self.assertEqual(comments[second["id"]]["anchorQuote"], "表格甲乙丙")
            edit_comment(reopened, {"id": first["id"], "expectedText": "甲乙丙", "newText": "甲丙"})
            reopened.save(filename)
            final = Document(filename)
            self.assertEqual(final.paragraphs[0].text, "前甲丙后")
            self.assertEqual(len(final.comments), 2)
            with self.assertRaisesRegex(ValueError, "锚点已失效"):
                edit_comment(final, {"id": first["id"], "expectedText": "甲乙丙", "newText": "错"})


if __name__ == "__main__":
    unittest.main()
