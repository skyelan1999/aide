#!/usr/bin/env python3
"""extract_text <path>：把 Office/PDF 文档提取为纯文本，打到 stdout。

供 aide 把 .docx/.pdf/.xlsx/.pptx 作为附件时，将正文随消息发给模型
（不依赖模型视觉能力）。失败时 stderr 报错并 exit 1。

分派：
  .docx  → python-docx（段落 + 表格）
  .xlsx  → openpyxl（按 sheet 逐行，单元格以 tab 分隔）
  .pptx  → python-pptx（逐幻灯片文本框）
  .pdf   → pypdf（逐页 extract_text）

按字符数和文本块数量在提取时截断；Go 层保留额外的输出预算。
"""
import os
import sys


# Bound output while iterating, before building one giant result string.
# 240000 Unicode characters leave room for UTF-8 expansion under Go's 1MiB cap.
MAX_TEXT_CHARS = 240000
MAX_TEXT_BLOCKS = 4096


def collect_text(lines):
    parts = []
    chars = 0
    truncated = False
    try:
        for line in lines:
            text = str(line or '').strip()
            if not text:
                continue
            remaining = MAX_TEXT_CHARS - chars
            if remaining <= 0 or len(parts) >= MAX_TEXT_BLOCKS:
                truncated = True
                break
            if len(text) > remaining:
                parts.append(text[:remaining])
                truncated = True
                break
            parts.append(text)
            chars += len(text) + 1
    finally:
        close = getattr(lines, 'close', None)
        if close:
            close()
    if truncated:
        parts.append('…[truncated]')
    return '\n'.join(parts)


def extract_docx(path: str) -> str:
    from docx import Document

    def lines():
        doc = Document(path)
        for paragraph in doc.paragraphs:
            yield paragraph.text or ''
        for index, table in enumerate(doc.tables):
            yield f"[表格 {index + 1}]"
            for row in table.rows:
                cells = [(c.text or '').strip().replace('\n', ' ') for c in row.cells]
                if any(cells):
                    yield '\t'.join(cells)
    return collect_text(lines())


def extract_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    def lines():
        wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in wb.worksheets:
                yield f"[Sheet: {sheet.title}]"
                for row in sheet.iter_rows(values_only=True):
                    cells = ['' if value is None else str(value).strip() for value in row]
                    while cells and not cells[-1]:
                        cells.pop()
                    if any(cells):
                        yield '\t'.join(cells)
        finally:
            wb.close()
    return collect_text(lines())


def extract_pptx(path: str) -> str:
    from pptx import Presentation

    def lines():
        for index, slide in enumerate(Presentation(path).slides):
            yield f"[幻灯片 {index + 1}]"
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        yield ''.join(run.text for run in paragraph.runs)
    return collect_text(lines())


def extract_pdf(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        print('容器未安装 pypdf，无法提取 PDF 文本', file=sys.stderr)
        return ''

    def lines():
        reader = PdfReader(path)
        for index, page in enumerate(reader.pages):
            text = (page.extract_text() or '').strip()
            if text:
                yield f"[第 {index + 1} 页]\n{text}"
    return collect_text(lines())


DISPATCH = {
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".pptx": extract_pptx,
    ".pdf": extract_pdf,
}


def main() -> int:
    if len(sys.argv) < 2:
        print("用法: extract_text.py <path>", file=sys.stderr)
        return 1
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    fn = DISPATCH.get(ext)
    if fn is None:
        print(f"不支持的文档类型: {ext}", file=sys.stderr)
        return 1
    if not os.path.isfile(path):
        print(f"文件不存在: {path}", file=sys.stderr)
        return 1
    try:
        text = fn(path)
    except Exception as e:  # 提取库任何异常都明确上报，不静默丢内容
        print(f"提取 {ext} 失败: {type(e).__name__}: {e}", file=sys.stderr)
        return 1
    if not text.strip():
        print(f"未能从 {ext} 提取到文本（可能是扫描件/图片型文档）: {path}", file=sys.stderr)
        return 1
    sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
