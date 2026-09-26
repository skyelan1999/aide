#!/usr/bin/env python3
"""docx_structure <path>：输出 .docx 的结构 JSON（标题层级/段落摘要/表格行列/批注数）。

供 aide 先"了解文档"再读批注。stdout 打印 JSON，失败时 stderr 报错并 exit 1。
"""
import json
import os
import sys


def main() -> int:
    path = sys.argv[1]
    if not path.lower().endswith(".docx"):
        print("仅支持 .docx 文件", file=sys.stderr)
        return 1
    if not os.path.isfile(path):
        print(f"文件不存在: {path}", file=sys.stderr)
        return 1

    from docx import Document  # noqa: WPS433（延迟导入，错误信息更清晰）

    d = Document(path)
    out = {"headings": [], "paragraphs": [], "tables": [], "commentCount": len(d.comments)}

    for p in d.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        style = p.style.name if p.style is not None else ""
        level = 0
        low = style.lower()
        if low.startswith("heading"):
            try:
                level = int(low.split()[-1])
            except ValueError:
                level = 1
        elif low.startswith("标题") or low.startswith("title"):
            level = 1 if low.startswith("title") else 2
        if level:
            out["headings"].append({"level": level, "text": text[:120]})
        else:
            out["paragraphs"].append({"index": len(out["paragraphs"]), "text": text[:50]})

    for i, t in enumerate(d.tables):
        out["tables"].append({"index": i, "rows": len(t.rows), "cols": len(t.columns)})

    print(json.dumps(out, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
