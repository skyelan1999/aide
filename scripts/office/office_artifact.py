#!/usr/bin/env python3
"""Bounded, structured Office creation and XLSX cell editing for aide.

The Go server owns paths and permissions. This helper only reads local temporary
inputs and writes a local temporary output; it never resolves workspace paths.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation


def fail(message):
    raise ValueError(message)


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def create(kind, spec, output):
    if kind == "docx":
        doc = Document()
        if spec.get("title"):
            doc.add_heading(str(spec["title"]), 0)
        for block in spec.get("blocks", []):
            if not isinstance(block, dict):
                fail("blocks must contain objects")
            block_type = block.get("type", "paragraph")
            if block_type == "heading":
                doc.add_heading(str(block.get("text", "")), level=max(1, min(9, int(block.get("level", 1)))))
            elif block_type == "paragraph":
                doc.add_paragraph(str(block.get("text", "")))
            elif block_type == "table":
                rows = block.get("rows", [])
                if not rows or not isinstance(rows, list) or not all(isinstance(row, list) for row in rows):
                    fail("table rows must be a non-empty list of lists")
                columns = max(map(len, rows))
                table = doc.add_table(rows=len(rows), cols=columns)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, value in enumerate(row):
                        table.cell(ri, ci).text = "" if value is None else str(value)
            else:
                fail("unsupported DOCX block type")
        doc.save(output)
    elif kind == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        sheets = spec.get("sheets", [])
        if not isinstance(sheets, list) or not sheets:
            fail("XLSX requires at least one sheet")
        for item in sheets:
            if not isinstance(item, dict):
                fail("sheet must be an object")
            ws = wb.create_sheet(str(item.get("name", "Sheet")))
            for row in item.get("rows", []):
                if not isinstance(row, list):
                    fail("sheet rows must contain lists")
                ws.append(row)
        wb.save(output)
    elif kind == "pptx":
        prs = Presentation()
        for item in spec.get("slides", []):
            if not isinstance(item, dict):
                fail("slide must be an object")
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = str(item.get("title", ""))
            slide.placeholders[1].text = str(item.get("body", ""))
        if not prs.slides:
            fail("PPTX requires at least one slide")
        prs.save(output)
    else:
        fail("unsupported Office format")


def view_xlsx(source, sheet_name, start_row, start_col):
    wb = load_workbook(source, read_only=True, data_only=False)
    try:
        if sheet_name and sheet_name not in wb.sheetnames:
            fail("sheet not found")
        ws = wb[sheet_name or wb.sheetnames[0]]
        rows = []
        for row in ws.iter_rows(min_row=start_row, max_row=min(start_row + 99, ws.max_row),
                                min_col=start_col, max_col=min(start_col + 25, ws.max_column)):
            rows.append([{"ref": cell.coordinate, "value": cell.value if cell.value is not None else ""}
                         for cell in row])
        return {"sheets": wb.sheetnames, "sheet": ws.title, "rows": rows,
                "maxRow": ws.max_row, "maxCol": ws.max_column,
                "startRow": start_row, "startCol": start_col}
    finally:
        wb.close()


def edit_xlsx(source, spec, output):
    wb = load_workbook(source, read_only=False, keep_links=True)
    try:
        changes = spec.get("changes", [])
        if not isinstance(changes, list) or not changes or len(changes) > 500:
            fail("changes must contain 1-500 cells")
        for change in changes:
            sheet = change.get("sheet")
            ref = change.get("ref")
            if sheet not in wb.sheetnames or not isinstance(ref, str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", ref):
                fail("invalid sheet or cell reference")
            cell = wb[sheet][ref]
            if cell.__class__.__name__ == "MergedCell":
                fail("cannot edit merged-cell continuation")
            value = change.get("value")
            if value is not None and not isinstance(value, (str, int, float, bool)):
                fail("unsupported cell value")
            cell.value = value
        wb.save(output)
    finally:
        wb.close()


def main(argv):
    if len(argv) < 2:
        fail("usage: office_artifact.py create|xlsx-view|xlsx-edit ...")
    op = argv[1]
    if op == "create" and len(argv) == 5:
        create(argv[2], read_json(argv[3]), argv[4])
    elif op == "xlsx-view" and len(argv) == 6:
        result = view_xlsx(argv[2], argv[3], int(argv[4]), int(argv[5]))
        print(json.dumps(result, ensure_ascii=False, default=str))
    elif op == "xlsx-edit" and len(argv) == 5:
        edit_xlsx(argv[2], read_json(argv[3]), argv[4])
    else:
        fail("invalid operation")


if __name__ == "__main__":
    try:
        main(sys.argv)
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)
