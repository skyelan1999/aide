"""Bounded source extraction. No document code, macros or external links execute."""
import json, sys
if sys.platform.startswith("linux"):
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (512 << 20, 512 << 20))
LIMIT = 256000
blocks = []; chars = 0; truncated = False

def add(locator, text):
    global chars, truncated
    text = str(text or '')
    if not text.strip(): return
    room = LIMIT - chars
    if room <= 0 or len(blocks) >= 4096:
        truncated = True; return
    if len(text) > room: text = text[:room]; truncated = True
    blocks.append({'locator': locator, 'text': text}); chars += len(text)

p, ext = sys.argv[1:3]
if ext == '.pdf':
    from pypdf import PdfReader
    reader = PdfReader(p)
    for i, page in enumerate(reader.pages):
        if truncated: break
        add('第 %d 页' % (i+1), page.extract_text() or '')
elif ext == '.docx':
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    d = Document(p); para = table = 0
    for child in d.element.body:
        if truncated: break
        if child.tag.endswith('}p'):
            para += 1; add('段落 %d' % para, Paragraph(child, d).text)
        elif child.tag.endswith('}tbl'):
            table += 1
            for i, row in enumerate(Table(child, d).rows):
                if truncated: break
                add('表格 %d / 行 %d' % (table, i+1), '\t'.join(c.text for c in row.cells))
elif ext == '.xlsx':
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=False, keep_links=False)
    for sheet in wb:
        if truncated: break
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if truncated: break
            add('工作表 %s / 行 %d' % (sheet.title, i+1), '\t'.join('' if x is None else str(x) for x in row))
    wb.close()
elif ext == '.pptx':
    from pptx import Presentation
    for i, slide in enumerate(Presentation(p).slides):
        if truncated: break
        for j, shape in enumerate(slide.shapes):
            if truncated: break
            if shape.has_text_frame: add('幻灯片 %d / 文本框 %d' % (i+1, j+1), shape.text)
            if shape.has_table:
                for k, row in enumerate(shape.table.rows):
                    if truncated: break
                    add('幻灯片 %d / 表格 %d / 行 %d' % (i+1, j+1, k+1), '\t'.join(c.text for c in row.cells))
else: raise ValueError('unsupported document format')
json.dump({'blocks': blocks, 'truncated': truncated}, sys.stdout, ensure_ascii=False)
